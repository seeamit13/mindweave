"""
Source ingestion: turn a PDF / PPTX / TXT file or a web URL into a list of
text chunks ready for embedding.

Two chunking strategies are implemented because the RAG evaluation
(section 6 of the assignment) compares chunking strategies as one of its
two retrieval approaches:

  - chunk_fixed():     naive fixed-size sliding-window chunking on raw characters
  - chunk_sentences(): sentence-aware chunking that packs whole sentences into
                        a chunk up to CHUNK_SIZE, which tends to keep semantic
                        units intact and improves retrieval precision.
"""

from __future__ import annotations

import re
from pathlib import Path
from dataclasses import dataclass

import requests
from bs4 import BeautifulSoup
from pypdf import PdfReader
from pptx import Presentation

from . import config


class IngestionError(Exception):
    """Raised when a source cannot be read/extracted."""


@dataclass
class Chunk:
    text: str
    chunk_index: int


# ---------------------------------------------------------------------------
# Text extraction
# ---------------------------------------------------------------------------

def extract_text_from_pdf(file_path: str) -> str:
    try:
        reader = PdfReader(file_path)
        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
        text = "\n\n".join(pages).strip()
        if not text:
            raise IngestionError("No extractable text found in PDF (it may be a scanned/image-only PDF).")
        return text
    except IngestionError:
        raise
    except Exception as e:
        raise IngestionError(f"Failed to read PDF: {e}") from e


def extract_text_from_pptx(file_path: str) -> str:
    try:
        prs = Presentation(file_path)
        slides_text = []
        for i, slide in enumerate(prs.slides, start=1):
            parts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        line = "".join(run.text for run in para.runs)
                        if line.strip():
                            parts.append(line)
                if shape.has_table:
                    for row in shape.table.rows:
                        parts.append(" | ".join(cell.text for cell in row.cells))
            if parts:
                slides_text.append(f"[Slide {i}]\n" + "\n".join(parts))
        text = "\n\n".join(slides_text).strip()
        if not text:
            raise IngestionError("No extractable text found in PPTX.")
        return text
    except IngestionError:
        raise
    except Exception as e:
        raise IngestionError(f"Failed to read PPTX: {e}") from e


def extract_text_from_txt(file_path: str) -> str:
    try:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            text = f.read().strip()
        if not text:
            raise IngestionError("TXT file is empty.")
        return text
    except IngestionError:
        raise
    except Exception as e:
        raise IngestionError(f"Failed to read TXT file: {e}") from e


def extract_text_from_url(url: str) -> str:
    try:
        resp = requests.get(
            url,
            timeout=config.REQUEST_TIMEOUT_SECS,
            headers={"User-Agent": "Mozilla/5.0 (NotebookCloneBot/1.0)"},
        )
        resp.raise_for_status()
    except requests.RequestException as e:
        raise IngestionError(f"Failed to fetch URL: {e}") from e

    soup = BeautifulSoup(resp.text, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header", "noscript"]):
        tag.decompose()

    text = soup.get_text(separator="\n")
    text = re.sub(r"\n{2,}", "\n\n", text).strip()
    text = text[: config.MAX_URL_FETCH_CHARS]

    if not text:
        raise IngestionError("No extractable text found at that URL.")
    return text


def extract_text(source_type: str, path_or_url: str) -> str:
    """Dispatch to the right extractor based on source_type."""
    source_type = source_type.lower()
    if source_type == "pdf":
        return extract_text_from_pdf(path_or_url)
    if source_type == "pptx":
        return extract_text_from_pptx(path_or_url)
    if source_type == "txt":
        return extract_text_from_txt(path_or_url)
    if source_type == "url":
        return extract_text_from_url(path_or_url)
    raise IngestionError(f"Unsupported source type: {source_type}")


def infer_source_type(filename: str) -> str:
    ext = Path(filename).suffix.lower().lstrip(".")
    if ext in ("pdf",):
        return "pdf"
    if ext in ("pptx",):
        return "pptx"
    if ext in ("txt", "md"):
        return "txt"
    raise IngestionError(f"Unsupported file type: .{ext}. Supported: .pdf, .pptx, .txt")


# ---------------------------------------------------------------------------
# Chunking strategies
# ---------------------------------------------------------------------------

def chunk_fixed(text: str, chunk_size: int = None, overlap: int = None) -> list[Chunk]:
    """Naive fixed-size sliding-window chunking over raw characters.
    Fast and simple, but can split sentences/words mid-way."""
    chunk_size = chunk_size or config.CHUNK_SIZE
    overlap = overlap if overlap is not None else config.CHUNK_OVERLAP
    text = text.strip()
    if not text:
        return []

    chunks = []
    start = 0
    idx = 0
    n = len(text)
    while start < n:
        end = min(start + chunk_size, n)
        piece = text[start:end].strip()
        if piece:
            chunks.append(Chunk(text=piece, chunk_index=idx))
            idx += 1
        if end == n:
            break
        start = end - overlap if end - overlap > start else end
    return chunks


_SENTENCE_SPLIT_RE = re.compile(r"(?<=[.!?])\s+")


def chunk_sentences(text: str, chunk_size: int = None, overlap: int = None) -> list[Chunk]:
    """Sentence-aware chunking: pack whole sentences into a chunk until adding
    another sentence would exceed chunk_size, then start a new chunk carrying
    a small sentence-level overlap forward. Keeps semantic units intact,
    which tends to improve retrieval precision vs. fixed-size chunking."""
    chunk_size = chunk_size or config.CHUNK_SIZE
    overlap = overlap if overlap is not None else config.CHUNK_OVERLAP
    text = text.strip()
    if not text:
        return []

    # Split on paragraphs first, then sentences, to preserve some structure.
    paragraphs = [p.strip() for p in re.split(r"\n{2,}", text) if p.strip()]
    sentences = []
    for para in paragraphs:
        sentences.extend(s for s in _SENTENCE_SPLIT_RE.split(para) if s.strip())

    chunks = []
    current = []
    current_len = 0
    idx = 0

    def flush():
        nonlocal current, current_len, idx
        if current:
            chunks.append(Chunk(text=" ".join(current).strip(), chunk_index=idx))
            idx += 1

    for sent in sentences:
        if current_len + len(sent) + 1 > chunk_size and current:
            flush()
            # carry overlap forward: keep trailing sentences up to `overlap` chars
            carried, carried_len = [], 0
            for s in reversed(current):
                if carried_len + len(s) > overlap:
                    break
                carried.insert(0, s)
                carried_len += len(s)
            current = carried
            current_len = carried_len
        current.append(sent)
        current_len += len(sent) + 1

    flush()
    return [c for c in chunks if c.text]


CHUNKING_STRATEGIES = {
    "fixed": chunk_fixed,
    "sentence": chunk_sentences,
}
